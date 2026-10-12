#!/usr/bin/env python3
"""生成层：把筛选+翻译后的条目渲染成 docx 简报，格式与样例一致。
输入: brief_content.json  [{source, zh_title, author, link, date, paragraphs:[...]}]
用法: python generate.py output/brief_content.json [输出.docx]"""
import json, sys, time
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

def build(items, out_path, brief_title=None):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    if brief_title:
        doc.add_heading(brief_title, level=1)

    for it in items:
        head = f"【{it['source']}】{it['zh_title']}"
        if it.get('author'):
            head += f"（作者：{it['author']}）"
        p = doc.add_paragraph()
        run = p.add_run(head); run.bold = True
        doc.add_paragraph('信源链接：')
        doc.add_paragraph(it['link'])
        prefix = f"{it['source']} {it['date']}报道 " if it.get('date') else ''
        paras = it['paragraphs']
        for i, para in enumerate(paras):
            doc.add_paragraph((prefix if i == 0 else '') + para)
        doc.add_paragraph('')  # 条目间隔
    doc.save(out_path)
    return out_path

def main():
    content_path = sys.argv[1] if len(sys.argv) > 1 else 'output/brief_content.json'
    out = sys.argv[2] if len(sys.argv) > 2 else f"output/简报_{time.strftime('%Y%m%d_%H%M')}.docx"
    items = json.load(open(content_path, encoding='utf-8'))
    path = build(items, out, brief_title=time.strftime('全球资讯简报 %Y-%m-%d %H:%M'))
    print('saved:', path, f'({len(items)} 条)')

if __name__ == '__main__':
    main()
